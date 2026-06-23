"""
summarize_matches.py
====================
NOAA / SWPC 매칭 결과 CSV 를 한 표로 통합·정렬해 비교하는 도구.
detector(GK2A / MetOp03 / NOAA19) × catalog(noaa / swpc) × method × channel
전체를 단일 마스터 표로 취합한다.

── glob 자동수집 ────────────────────────────────────────────────────
--dir 하나로 noaa_match_*.csv / swpc_match_*.csv 를 재귀 탐색.
카탈로그(noaa/swpc)는 파일명 접두에서 자동 판별.
detector 는 파일 경로(MetOp03_count / NOAA19_count → metop03/noaa19,
그 외 → gk2a)에서 자동 판별.

── J' 정의 ─────────────────────────────────────────────────────────
J' = POD - FAR.

사용:
  # 디렉터리 재귀 (noaa + swpc 모두)
  python summarize_matches.py --dir /path/to/results --out match_consolidated

  # 특정 파일들만
  python summarize_matches.py --inputs a.csv b.csv --out cmp

  # 상위 N 채널
  python summarize_matches.py --dir . --topk 3
"""
from __future__ import annotations
import re
import sys
import argparse
from pathlib import Path

import numpy as np
import pandas as pd

BIN_COLS = ["POD_10-100", "POD_100-1k", "POD_1k+"]

CATALOG_PREFIXES = ["noaa", "swpc"]   # 인식 가능한 파일 접두


# ── J' ──────────────────────────────────────────────────────────────
def _jprime(df: pd.DataFrame) -> pd.Series:
    return df["POD"] - df["FAR"]


# ── 채널 분류 ────────────────────────────────────────────────────────
def _channel_type(channel: str) -> str:
    """GK2A: 'PD3A-OU' 형식 → 'geo'.  POES: 'ele_tel0_e1' 형식 → 'poes'."""
    return "geo" if "-" in channel else "poes"


def _split_geo_channel(channel: str):
    """'PD3A-OU' → (pd_key='PD3', side='A', logic='OU')."""
    try:
        pd_side, logic = channel.split("-", 1)
        return pd_side[:-1], pd_side[-1], logic
    except ValueError:
        return np.nan, np.nan, np.nan


def _split_poes_channel(channel: str):
    """'ele_tel0_e1' → (species='ele', direction='tel0', energy='e1')."""
    parts = channel.split("_", 2)
    if len(parts) == 3:
        return parts[0], parts[1], parts[2]
    return np.nan, np.nan, np.nan


def _group_geo(logic: str) -> str:
    if logic in ("OU", "OUT"):   return "A"
    if logic in ("FTU", "FTUO"): return "B"
    return "C"


def _group_poes(species: str) -> str:
    """ele → A (electron = SWPC 적합), pro → B (NOAA 적합), omni → C."""
    if species == "ele":  return "A"
    if species == "pro":  return "B"
    return "C"


# ── 파일명 파싱 ──────────────────────────────────────────────────────
def _parse_catalog(filename: str) -> str:
    """'noaa_match_...csv' → 'noaa', 'swpc_match_...csv' → 'swpc'."""
    for pfx in CATALOG_PREFIXES:
        if filename.startswith(f"{pfx}_match_"):
            return pfx
    return "unknown"


def _parse_method(path: Path) -> str:
    """파일명에서 방법 식별자 추출. summary_all / onset / event 접미 제거."""
    stem = path.stem
    for pfx in CATALOG_PREFIXES:
        stem = re.sub(rf"^{pfx}_match_", "", stem)
    stem = re.sub(r"^summary_all_", "", stem)
    stem = re.sub(r"_(onset|nopk|event)$", "", stem)
    return stem


def _method_family(method: str) -> str:
    return re.split(r"_w\d|_on", method)[0]


def _detect_detector(path: Path) -> str:
    """파일 경로에서 detector 판별.
    MetOp03_count → metop03, NOAA19_count → noaa19, 그 외 → gk2a."""
    parts = [p.lower() for p in path.parts]
    if any("metop03" in p for p in parts): return "metop03"
    if any("noaa19"  in p for p in parts): return "noaa19"
    return "gk2a"


# ── 파일 수집 ────────────────────────────────────────────────────────
def discover(dir_path: Path, use_summary: bool) -> list[Path]:
    files = []
    for pfx in CATALOG_PREFIXES:
        files.extend(dir_path.rglob(f"{pfx}_match_*.csv"))
    files = sorted(set(files))
    if not use_summary:
        files = [f for f in files if "summary_all" not in f.name]
    return files


def load_all(files: list[Path], use_summary: bool) -> pd.DataFrame:
    seen, frames = {}, []
    for f in files:
        catalog  = _parse_catalog(f.name)
        method   = _parse_method(f)
        detector = _detect_detector(f)
        key      = (catalog, detector, method)

        is_summary = "summary_all" in f.name
        if key in seen:
            if seen[key] == "summary" and not is_summary:
                frames = [fr for fr in frames
                          if not ((fr["catalog"].iloc[0] == catalog) and
                                  (fr["detector"].iloc[0] == detector) and
                                  (fr["method"].iloc[0] == method))]
            else:
                continue
        seen[key] = "summary" if is_summary else "final"

        df = pd.read_csv(f)
        if df.empty or "channel" not in df.columns:
            print(f"[skip] {f.name}: 빈 결과 — 제외")
            continue

        df["catalog"]       = catalog
        df["detector"]      = detector
        df["method"]        = method
        df["method_family"] = _method_family(method)
        df["Jprime"]        = _jprime(df)
        df["source_file"]   = f.name

        # 채널 분해 (GEO / POES 모두 처리)
        ctype = df["channel"].iloc[0]
        if "-" in ctype:   # GEO
            pds = df["channel"].apply(_split_geo_channel)
            df["ch_type"] = "geo"
            df["ch_a"]    = [p[0] for p in pds]   # pd_key
            df["ch_b"]    = [p[1] for p in pds]   # side
            df["ch_c"]    = [p[2] for p in pds]   # logic
            df["group"]   = df["ch_c"].map(_group_geo)
        else:              # POES
            pds = df["channel"].apply(_split_poes_channel)
            df["ch_type"] = "poes"
            df["ch_a"]    = [p[0] for p in pds]   # species
            df["ch_b"]    = [p[1] for p in pds]   # direction
            df["ch_c"]    = [p[2] for p in pds]   # energy
            df["group"]   = df["ch_a"].map(_group_poes)

        frames.append(df)

    if not frames:
        raise SystemExit("[summarize] 읽을 매칭 CSV가 없습니다.")
    return pd.concat(frames, ignore_index=True)


# ── 뷰 생성 ─────────────────────────────────────────────────────────
LEAD_COLS = [
    "catalog", "detector", "method_family", "method",
    "channel", "ch_type", "ch_a", "ch_b", "ch_c", "group",
    "POD", "FAR", "Jprime",
    "n_det", "n_hit", "n_fa",
    "onset_diff_med_h", "peak_diff_med_h",
    *BIN_COLS,
]


def view_all(comb: pd.DataFrame) -> pd.DataFrame:
    cols = [c for c in LEAD_COLS if c in comb.columns]
    return comb[cols].sort_values(
        ["catalog", "detector", "method_family", "method", "Jprime"],
        ascending=[True, True, True, True, False]
    ).reset_index(drop=True)


def view_best_per_method(comb: pd.DataFrame, topk: int) -> pd.DataFrame:
    cols = ["catalog", "detector", "method", "channel",
            "ch_c", "group", "POD", "FAR", "Jprime",
            "n_det", "n_hit", "n_fa"]
    cols = [c for c in cols if c in comb.columns]
    return (comb.sort_values("Jprime", ascending=False)
                .groupby(["catalog", "detector", "method"], sort=False)
                .head(topk)[cols]
                .sort_values(["catalog", "detector", "method", "Jprime"],
                             ascending=[True, True, True, False])
                .reset_index(drop=True))


def _agg_by(comb: pd.DataFrame, key: str) -> pd.DataFrame:
    rows = []
    for (catalog, detector, method, kv), g in comb.groupby(
            ["catalog", "detector", "method", key]):
        gg   = g.dropna(subset=["Jprime"])
        best = gg.loc[gg["Jprime"].idxmax()] if len(gg) else g.iloc[0]
        rows.append({
            "catalog": catalog, "detector": detector, "method": method, key: kv,
            "n_ch":         len(g),
            "POD_mean":     round(g["POD"].mean(), 3),
            "FAR_mean":     round(g["FAR"].mean(), 3),
            "Jprime_mean":  round(g["Jprime"].mean(), 3),
            "Jprime_best":  round(float(best["Jprime"]), 3) if len(gg) else np.nan,
            "best_channel": best["channel"] if len(gg) else "",
            "POD_best":     round(float(best["POD"]), 3) if len(gg) else np.nan,
            "FAR_best":     round(float(best["FAR"]), 3) if len(gg) else np.nan,
        })
    return pd.DataFrame(rows).sort_values(
        ["catalog", "detector", "method", "Jprime_mean"],
        ascending=[True, True, True, False]).reset_index(drop=True)


def view_by_group(comb): return _agg_by(comb, "group")
def view_by_ch_c(comb):  return _agg_by(comb, "ch_c")   # logic (GEO) / energy (POES)


def view_method_rank(comb: pd.DataFrame) -> pd.DataFrame:
    """(catalog × detector × method) 랭킹. group A 평균 J' 기준."""
    rows = []
    for (catalog, detector, method), g in comb.groupby(
            ["catalog", "detector", "method"]):
        a    = g[g["group"] == "A"].dropna(subset=["Jprime"])
        gg   = g.dropna(subset=["Jprime"])
        best = gg.loc[gg["Jprime"].idxmax()] if len(gg) else None
        rows.append({
            "catalog":        catalog,
            "detector":       detector,
            "method":         method,
            "family":         g["method_family"].iloc[0],
            "JprimeA_mean":   round(a["Jprime"].mean(), 3) if len(a) else np.nan,
            "PODA_mean":      round(a["POD"].mean(),    3) if len(a) else np.nan,
            "FARA_mean":      round(a["FAR"].mean(),    3) if len(a) else np.nan,
            "n_chA":          len(a),
            "Jprime_best_any": round(float(best["Jprime"]), 3) if best is not None else np.nan,
            "best_channel_any": best["channel"] if best is not None else "",
        })
    return pd.DataFrame(rows).sort_values(
        ["catalog", "detector", "JprimeA_mean"],
        ascending=[True, True, False]).reset_index(drop=True)


# ── xlsx ─────────────────────────────────────────────────────────────
def write_xlsx(views: dict, path: Path) -> bool:
    try:
        from openpyxl.styles import Font, Alignment, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return False
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for sheet, df in views.items():
            df.to_excel(xw, sheet_name=sheet[:31], index=False)
        wb = xw.book
        head_fill = PatternFill("solid", start_color="1F4E78")
        head_font = Font(bold=True, color="FFFFFF", name="Arial")
        for sheet, df in views.items():
            ws = wb[sheet[:31]]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for ci, col in enumerate(df.columns, 1):
                cell = ws.cell(row=1, column=ci)
                cell.fill = head_fill; cell.font = head_font
                cell.alignment = Alignment(horizontal="center")
                width = max(len(str(col)),
                            *(len(str(v)) for v in df[col].astype(str).head(50)), 6)
                ws.column_dimensions[get_column_letter(ci)].width = min(width + 2, 40)
    return True


# ── main ─────────────────────────────────────────────────────────────
def main():
    ap = argparse.ArgumentParser(
        description="NOAA+SWPC / GEO+LEO 매칭 결과 통합·비교 (J' 정렬, group 집계)")
    src = ap.add_mutually_exclusive_group(required=True)
    src.add_argument("--dir",    help="결과 루트 폴더(재귀 탐색, noaa+swpc 모두)")
    src.add_argument("--inputs", nargs="+", help="CSV 파일들 직접 지정")
    ap.add_argument("--out",   default="match_consolidated", help="출력 폴더")
    ap.add_argument("--topk",  type=int, default=5, help="방법별 상위 채널 수")
    ap.add_argument("--use-summary", action="store_true",
                    help="summary_all CSV도 포함(기본은 FINAL만)")
    args = ap.parse_args()

    if args.dir:
        files = discover(Path(args.dir), args.use_summary)
    else:
        files = [Path(p) for p in args.inputs]
    if not files:
        raise SystemExit("[summarize] 입력 CSV를 찾지 못했습니다.")
    print(f"[summarize] {len(files)}개 파일 통합:")
    for f in files:
        print(f"   {f}")

    comb = load_all(files, args.use_summary)
    print(f"[summarize] catalog×detector×method×channel → 총 {len(comb)}행\n")

    views = {
        "all_channels":    view_all(comb),
        "best_per_method": view_best_per_method(comb, args.topk),
        "by_group":        view_by_group(comb),
        "by_ch_c":         view_by_ch_c(comb),
        "method_rank":     view_method_rank(comb),
    }

    outdir = Path(args.out); outdir.mkdir(parents=True, exist_ok=True)
    for name, df in views.items():
        df.to_csv(outdir / f"{name}.csv", index=False)
    xlsx_ok = write_xlsx(views, outdir / "match_comparison.xlsx")

    print("[summarize] 저장:")
    for name in views:
        print(f"   {outdir / (name + '.csv')}")
    if xlsx_ok:
        print(f"   {outdir / 'match_comparison.xlsx'}  (멀티시트)")
    else:
        print("   (openpyxl 없음 → xlsx 생략, CSV만)")

    pd.set_option("display.width", 180, "display.max_columns", 30)
    print("\n===== 방법 랭킹 (catalog×detector, group A 평균 J') =====")
    print(views["method_rank"].to_string(index=False))
    print("\n===== 방법별 최고 채널 (J' top) =====")
    print(views["best_per_method"].groupby(
        ["catalog", "detector", "method"], sort=False).head(3).to_string(index=False))


if __name__ == "__main__":
    main()
